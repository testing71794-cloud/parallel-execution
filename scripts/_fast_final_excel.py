#!/usr/bin/env python3
"""Fast final Excel from status/*.txt + log tails (no ADB / OpenRouter)."""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parents[1]
STATUS = REPO / "status"
OUT = REPO / "build-summary" / "final_execution_report.xlsx"

DEVICE_NAMES = {
    "16091FDD4004N6": "Pixel 5",
    "ZA222RFQ75": "Motorola Fusion 50",
}

PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
HEADER_FILL = PatternFill("solid", fgColor="2E5C8A")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def parse_status(path: Path) -> dict[str, str]:
    data: dict[str, str] = {}
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return data
    for line in text.splitlines():
        if "=" in line:
            k, v = line.split("=", 1)
            data[k.strip().lower()] = v.strip()
    return data


def reason_from_log(log_path: str, fallback: str) -> str:
    try:
        from status_normalize import extract_failure_reason_from_log

        return extract_failure_reason_from_log(log_path, fallback=fallback or "MAESTRO_FAILED")
    except Exception:
        pass
    p = Path(log_path) if log_path else None
    if not p or not p.is_file():
        return fallback or "Unknown failure"
    try:
        raw = p.read_bytes()
        if len(raw) > 200_000:
            raw = raw[-200_000:]
        text = raw.decode("utf-8", errors="replace")
    except OSError:
        return fallback or "Unknown failure"

    patterns = [
        r"Assertion is false:.*",
        r"Element not found:.*",
        r"Unable to find element.*",
        r"INSTALL_FAILED_[A-Z_]+.*",
        r"Device .* is not connected.*",
        r"Assertion failed.*",
        r"Timeout.*",
        r"Final reason\s*:\s*(.+)",
        r"FAILURE:.*",
        r"Error:.*",
    ]
    for pat in patterns:
        m = re.search(pat, text, re.IGNORECASE | re.MULTILINE)
        if m:
            line = m.group(0).strip()
            return re.sub(r"\s+", " ", line)[:500]
    # last non-empty Maestro-ish line
    for line in reversed(text.splitlines()):
        s = line.strip()
        if s and not s.startswith("[") and len(s) > 20:
            return s[:500]
    return fallback or "MAESTRO_FAILED"


def latest_status_rows() -> list[dict]:
    best: dict[tuple[str, str, str], tuple[float, dict]] = {}
    for path in STATUS.glob("*.txt"):
        segs = path.stem.split("__")
        if len(segs) < 3:
            continue
        suite, device = segs[0], segs[-1]
        flow = "__".join(segs[1:-1])
        fields = parse_status(path)
        status_raw = (fields.get("status") or fields.get("final_status") or "").upper()
        reason_code = fields.get("reason") or fields.get("final_reason") or ""
        try:
            from status_normalize import normalize_status

            status = normalize_status(status_raw, reason_code)
        except Exception:
            status = "FAIL" if status_raw == "FLAKY" else status_raw
        if not status or status in ("RUNNING",):
            continue
        mtime = path.stat().st_mtime
        key = (suite, flow, device)
        row = {
            "Suite": suite,
            "Flow Name": flow.replace("_", " "),
            "Device Name": DEVICE_NAMES.get(device, device),
            "Device ID": device,
            "Status": status,
            "Exit Code": fields.get("exit_code") or fields.get("exit") or "",
            "Failure Reason": "",
            "Log Path": fields.get("log_file")
            or fields.get("log_path")
            or fields.get("log")
            or "",
            "Video Path": fields.get("video_file")
            or fields.get("video_path")
            or fields.get("recording_file")
            or "",
            "Timestamp": datetime.fromtimestamp(mtime).isoformat(timespec="seconds"),
        }
        if status != "PASS":
            row["Failure Reason"] = reason_from_log(
                row["Log Path"],
                fields.get("reason") or fields.get("final_reason") or "MAESTRO_FAILED",
            )
        prev = best.get(key)
        if prev is None or mtime >= prev[0]:
            best[key] = (mtime, row)
    rows = [v[1] for v in best.values()]
    rows.sort(key=lambda r: (r["Suite"], r["Flow Name"], r["Device Name"]))
    return rows


def main() -> int:
    rows = latest_status_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "Raw Results"
    headers = [
        "Suite",
        "Flow Name",
        "Device Name",
        "Device ID",
        "Status",
        "Exit Code",
        "Failure Reason",
        "Log Path",
        "Video Path",
        "Timestamp",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
        st = row["Status"].upper()
        cell = ws.cell(ws.max_row, 5)
        cell.fill = PASS_FILL if st == "PASS" else FAIL_FILL
        cell.font = Font(bold=True)

    fr = wb.create_sheet("Flow Report", 1)
    fr_headers = ["Suite", "Flow", "Device", "Device ID", "Status", "Exit Code", "Failure Reason"]
    fr.append(fr_headers)
    for c in fr[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
    for row in rows:
        fr.append(
            [
                row["Suite"],
                row["Flow Name"],
                row["Device Name"],
                row["Device ID"],
                row["Status"],
                row["Exit Code"],
                row["Failure Reason"] or ("—" if row["Status"] == "PASS" else "MAESTRO_FAILED"),
            ]
        )
        st = row["Status"].upper()
        cell = fr.cell(fr.max_row, 5)
        cell.fill = PASS_FILL if st == "PASS" else FAIL_FILL
        cell.font = Font(bold=True)

    # Summary
    sm = wb.create_sheet("Summary", 0)
    pass_n = sum(1 for r in rows if r["Status"] == "PASS")
    fail_n = sum(1 for r in rows if r["Status"] != "PASS")
    sm["A1"] = "Kodak Smile — Final Execution Report"
    sm["A1"].font = Font(bold=True, size=14)
    sm["A3"] = "Generated"
    sm["B3"] = datetime.now().isoformat(timespec="seconds")
    sm["A4"] = "Total rows"
    sm["B4"] = len(rows)
    sm["A5"] = "PASS"
    sm["B5"] = pass_n
    sm["A6"] = "FAIL/other"
    sm["B6"] = fail_n
    sm["A8"] = "Devices"
    sm["B8"] = "Pixel 5 (16091FDD4004N6), Motorola Fusion 50 (ZA222RFQ75)"
    sm["A9"] = "Failed videos zip"
    sm["B9"] = str(REPO / "build-summary" / "failed_tests_artifacts.zip")
    sm["A10"] = "Failed artifacts folder"
    sm["B10"] = str(REPO / "build-summary" / "failed-artifacts")

    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = min(48, max(12, len(str(col[0].value or "")) + 4))
    for col in fr.columns:
        letter = col[0].column_letter
        fr.column_dimensions[letter].width = min(56, max(12, len(str(col[0].value or "")) + 4))
        for cell in col[1:]:
            if letter == "G":
                cell.alignment = Alignment(wrap_text=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_bytes(b"")  # ensure writable
    wb.save(OUT)
    print(f"Wrote {OUT} rows={len(rows)} pass={pass_n} fail={fail_n}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
